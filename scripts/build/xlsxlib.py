"""xlsx 書式の記入・生成の共通ライブラリ。

2つのモードを支える:
1. テンプレートモード — 裁判所配布 xlsx（courts/ に登録済み）を fillmap 駆動で記入。
   数式・結合セル・印刷設定を保持するため、既存ブックを開いて値だけを書く。
2. 汎用ドラフトモード — 書式未登録のあいだの事務所内ドラフトを openpyxl で新規生成。
   builder スクリプトが章立てを組む。

fillmap（xlsx用）の kind:
  set_cell    anchor: {sheet, cell} に値を書く。label_check: {アドレス: 期待文言} で
              近傍ラベルを検証（不一致は警告＝書式改訂の兆候）
  table_rows  anchor: {sheet}, first_data_row, columns: {列レター: キー}, 値は辞書のリスト。
              max_rows を超えたら行挿入（直前行の書式をコピー）
  checkbox    セル値内の「□ラベル」→「☑ラベル」置換（空白ゆらぎ許容）。
              options: {値: リテラル}（anchor のセル内）または
              option_anchors: {値: {sheet, cell, literal}}
  row_blocks  複数行1組の繰返しブロック（1債権者=3行等）。
              anchor: {sheet}, first_row, block_rows, count,
              cells: [{r: 行オフセット0起点, c: 列レター, key: 値キー, checkbox: true(任意)}]
              値は辞書のリスト。count 超過分は書ききれず警告（呼び出し側で2枚目対応）

共通の防御: 既存セルが数式（"=" 始まり）の場合は上書きせず警告する。
"""
import re as _re
from pathlib import Path

import yaml
from copy import copy as style_copy

import openpyxl

REPO = Path(__file__).resolve().parent.parent.parent


def resolve_form(court: str, form: str):
    """registry.yaml から現行版のディレクトリを返す。未登録なら None。"""
    reg_path = REPO / "courts" / court / "forms" / form / "registry.yaml"
    if not reg_path.exists():
        return None
    reg = yaml.safe_load(reg_path.read_text(encoding="utf-8"))
    cur = reg.get("current")
    if not cur:
        return None
    vdir = reg_path.parent / cur
    if not (vdir / "fillmap.yaml").exists():
        return None
    return vdir


def registration_guidance(court: str, form: str, name: str) -> str:
    return (
        f"裁判所書式「{name}」が未登録です（courts/{court}/forms/{form}/）。\n"
        f"裁判所配布の Excel ファイルを入手したら hasan-register-form スキルで登録してください\n"
        f"（登録するまでは --generic で事務所内ドラフト様式を生成できます）。"
    )


class XlsxFiller:
    def __init__(self, template_path):
        self.wb = openpyxl.load_workbook(template_path)
        self.warnings = []
        self.filled = 0

    def _ws(self, name):
        if name in self.wb.sheetnames:
            return self.wb[name]
        self.warnings.append(f"シート「{name}」が無い（先頭シートで代替）")
        return self.wb.worksheets[0]

    def _write(self, ws, addr, value, fid):
        cur = ws[addr].value
        if isinstance(cur, str) and cur.startswith("="):
            self.warnings.append(f"{fid}: {addr} は数式セルのため書き込まない（{cur[:30]}）")
            return False
        ws[addr] = value
        return True

    def set_cell(self, field, value):
        ws = self._ws(field["anchor"]["sheet"])
        for addr, expect in (field.get("label_check") or {}).items():
            actual = ws[addr].value
            if actual != expect:
                self.warnings.append(
                    f"{field['id']}: ラベル検証不一致 {addr}=「{actual}」期待「{expect}」（書式改訂の疑い）")
        if self._write(ws, field["anchor"]["cell"], value, field["id"]):
            self.filled += 1

    @staticmethod
    def _flex_check(text, literal):
        """セル値文字列内の「□ラベル」を空白ゆらぎ許容で「☑ラベル」化。失敗は None。"""
        chars = [ch for ch in literal if not ch.isspace() and ch != "　"]
        pattern = "[\\s　]*".join(_re.escape(ch) for ch in chars)
        m = _re.search(pattern, text or "")
        if not m:
            return None
        seg = m.group(0)
        return (text[: m.start()] + seg.replace("□", "☑", 1) + text[m.end():])

    def checkbox(self, field, value):
        values = value if isinstance(value, list) else [value]
        for v in values:
            v = str(v)
            if "option_anchors" in field:
                spec = field["option_anchors"].get(v)
                if spec is None:
                    self.warnings.append(f"{field['id']}: 未知の選択肢「{v}」")
                    continue
                sheet, cell, literal = spec["sheet"], spec["cell"], spec["literal"]
            else:
                literal = (field.get("options") or {}).get(v)
                if literal is None:
                    self.warnings.append(f"{field['id']}: 未知の選択肢「{v}」")
                    continue
                sheet, cell = field["anchor"]["sheet"], field["anchor"]["cell"]
            ws = self._ws(sheet)
            new = self._flex_check(ws[cell].value, literal)
            if new is None:
                self.warnings.append(f"{field['id']}: {cell} にリテラル「{literal}」が見つからない")
            else:
                ws[cell] = new
                self.filled += 1

    def row_blocks(self, field, rows):
        ws = self._ws(field["anchor"]["sheet"])
        first = field["first_row"]
        block = field["block_rows"]
        count = field["count"]
        if len(rows) > count:
            self.warnings.append(
                f"{field['id']}: {len(rows)}件中{count}件のみ記入（書式の枠数超過。2枚目の作成が必要）")
        for i, row in enumerate(rows[:count]):
            base = first + i * block
            for cell_spec in field["cells"]:
                v = row.get(cell_spec["key"])
                if v is None or v == "":
                    continue
                addr = f"{cell_spec['c']}{base + cell_spec.get('r', 0)}"
                if cell_spec.get("checkbox"):
                    new = self._flex_check(ws[addr].value, str(v))
                    if new is None:
                        self.warnings.append(f"{field['id']}: {addr} に「{v}」のチェック対象が見つからない")
                    else:
                        ws[addr] = new
                else:
                    self._write(ws, addr, v, field["id"])
        self.filled += min(len(rows), count)

    def table_rows(self, field, rows):
        ws = self._ws(field["anchor"]["sheet"])
        first = field["first_data_row"]
        cols = field["columns"]
        max_rows = field.get("max_rows")
        if max_rows and len(rows) > max_rows:
            # 最終データ行の直後に必要数を挿入し、書式をコピー
            insert_at = first + max_rows
            n = len(rows) - max_rows
            ws.insert_rows(insert_at, n)
            for i in range(n):
                for col in cols:
                    src = ws[f"{col}{insert_at - 1}"]
                    dst = ws[f"{col}{insert_at + i}"]
                    dst.font = style_copy(src.font)
                    dst.border = style_copy(src.border)
                    dst.alignment = style_copy(src.alignment)
                    dst.number_format = src.number_format
        for i, row in enumerate(rows):
            for col, key in cols.items():
                v = row.get(key)
                if v is not None:
                    self._write(ws, f"{col}{first + i}", v, field["id"])
        self.filled += len(rows)

    def apply(self, fillmap, values):
        vals = values.get("fields", {})
        for field in fillmap["fields"]:
            v = vals.get(field["id"])
            if v is None or v == "" or v == []:
                continue
            if field["kind"] == "set_cell":
                self.set_cell(field, v)
            elif field["kind"] == "table_rows":
                self.table_rows(field, v)
            elif field["kind"] == "checkbox":
                self.checkbox(field, v)
            elif field["kind"] == "row_blocks":
                self.row_blocks(field, v)
            else:
                self.warnings.append(f"{field['id']}: 未知の kind {field['kind']}")

    def save(self, out):
        Path(out).parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out)


# ---- 汎用ドラフトモードの体裁ヘルパー ----------------------------------
THIN = openpyxl.styles.Side(style="thin")
BORDER = openpyxl.styles.Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = openpyxl.styles.PatternFill("solid", fgColor="EEEEEE")
FONT = openpyxl.styles.Font(name="游ゴシック", size=10)
FONT_TITLE = openpyxl.styles.Font(name="游ゴシック", size=14, bold=True)


def style_header(ws, row, ncols, start_col=1):
    for i in range(ncols):
        c = ws.cell(row=row, column=start_col + i)
        c.border = BORDER
        c.fill = HEAD_FILL
        c.font = openpyxl.styles.Font(name="游ゴシック", size=10, bold=True)
        c.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_data(ws, row, ncols, start_col=1):
    for i in range(ncols):
        c = ws.cell(row=row, column=start_col + i)
        c.border = BORDER
        c.font = FONT
        c.alignment = openpyxl.styles.Alignment(vertical="center", wrap_text=True)


def draft_note(ws, row, text="※ 事務所内ドラフト様式（裁判所書式は未登録）。提出前に裁判所書式へ転記・登録すること。"):
    c = ws.cell(row=row, column=1, value=text)
    c.font = openpyxl.styles.Font(name="游ゴシック", size=9, italic=True, color="AA0000")
