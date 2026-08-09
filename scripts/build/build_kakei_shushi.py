#!/usr/bin/env python3
"""家計収支表を case.yaml から生成する（計算ミスなく月次の収支を集計）。

裁判所書式（courts/<court>/forms/kakei-shushi/）が登録済みなら fillmap 駆動、
未登録なら --generic で事務所内ドラフト様式（月を列に並べた対比表）を出力する。
<output>.meta.json に月次集計を書き出し、crosscheck_outputs.py が突合に使う。

使い方:
  python3 scripts/build/build_kakei_shushi.py --case cases/<id>/case.yaml \
      --output cases/<id>/output/家計収支表_<姓>.xlsx [--court osaka] [--generic]
"""
import argparse
import json
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsxlib  # noqa: E402

import openpyxl  # noqa: E402


def month_stats(months):
    out = []
    for h in months:
        inc = sum(v for v in (h.get("income") or {}).values() if isinstance(v, int))
        exp = sum(v for v in (h.get("expense") or {}).values() if isinstance(v, int))
        out.append({"month": h.get("month", ""), "income": inc, "expense": exp, "balance": inc - exp})
    return out


def build_generic(case, out):
    meta = case.get("meta", {})
    months = case.get("household", []) or []
    if not months:
        print("household が空のため家計収支表を生成できない")
        sys.exit(4)

    # 費目の全集合（記載順を保ちつつ和集合）
    inc_keys, exp_keys = [], []
    for h in months:
        for k in (h.get("income") or {}):
            if k not in inc_keys:
                inc_keys.append(k)
        for k in (h.get("expense") or {}):
            if k not in exp_keys:
                exp_keys.append(k)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "家計収支表"
    ws["A1"] = "家計収支表"
    ws["A1"].font = xlsxlib.FONT_TITLE
    ws["A2"] = f"申立人: {meta.get('applicant', {}).get('name', '')}"
    ws["A2"].font = xlsxlib.FONT

    ncols = 1 + len(months)
    ws.cell(row=4, column=1, value="費目")
    ws.column_dimensions["A"].width = 18
    for j, h in enumerate(months, 2):
        ws.cell(row=4, column=j, value=h.get("month", ""))
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 14
    xlsxlib.style_header(ws, 4, ncols)

    def section(row, label, keys, kind):
        ws.cell(row=row, column=1, value=label)
        xlsxlib.style_header(ws, row, ncols)
        first = row + 1
        for k in keys:
            row += 1
            ws.cell(row=row, column=1, value=k)
            for j, h in enumerate(months, 2):
                v = (h.get(kind) or {}).get(k)
                ws.cell(row=row, column=j, value=v)
                ws.cell(row=row, column=j).number_format = "#,##0"
            xlsxlib.style_data(ws, row, ncols)
        row += 1
        ws.cell(row=row, column=1, value=f"{label}合計")
        for j in range(2, ncols + 1):
            col = openpyxl.utils.get_column_letter(j)
            ws.cell(row=row, column=j, value=f"=SUM({col}{first}:{col}{row - 1})")
            ws.cell(row=row, column=j).number_format = "#,##0"
        xlsxlib.style_data(ws, row, ncols)
        ws.cell(row=row, column=1).font = openpyxl.styles.Font(name="游ゴシック", size=10, bold=True)
        return row

    inc_total_row = section(5, "収入", inc_keys, "income")
    exp_total_row = section(inc_total_row + 1, "支出", exp_keys, "expense")

    row = exp_total_row + 1
    ws.cell(row=row, column=1, value="収支（収入−支出）")
    for j in range(2, ncols + 1):
        col = openpyxl.utils.get_column_letter(j)
        ws.cell(row=row, column=j, value=f"={col}{inc_total_row}-{col}{exp_total_row}")
        ws.cell(row=row, column=j).number_format = "#,##0;[Red]-#,##0"
    xlsxlib.style_data(ws, row, ncols)
    ws.cell(row=row, column=1).font = openpyxl.styles.Font(name="游ゴシック", size=10, bold=True)
    xlsxlib.draft_note(ws, row + 2)

    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"months": month_stats(months)}


def zenkaku(n):
    return str(n).translate(str.maketrans("0123456789", "０１２３４５６７８９"))


def month_period(m):
    """'R7.6' → '６月１日～６月３０日'（月末日は暦から計算）。形式外はそのまま返す。"""
    import calendar
    import re
    mt = re.match(r"^([MTSHR])(\d+)\.(\d+)$", m or "")
    if not mt:
        return m or ""
    base = {"M": 1867, "T": 1911, "S": 1925, "H": 1988, "R": 2018}[mt.group(1)]
    year, mon = base + int(mt.group(2)), int(mt.group(3))
    last = calendar.monthrange(year, mon)[1]
    return f"{zenkaku(mon)}月１日～{zenkaku(mon)}月{zenkaku(last)}日"


def title_year(m):
    import re
    mt = re.match(r"^R(\d+)\.", m or "")
    return f"　　　　　　　　　家計収支表(令和{zenkaku(int(mt.group(1)))}年）" if mt else ""


def build_from_template(case, vdir, out):
    """B1111: 費目固定行（fillmap item_rows）へ2か月分を記入。数式セルは触らない。"""
    fillmap = yaml.safe_load((vdir / "fillmap.yaml").read_text(encoding="utf-8"))
    months = (case.get("household", []) or [])[:2]
    if not months:
        print("household が空のため家計収支表を生成できない")
        sys.exit(4)
    if len(case.get("household", [])) > 2:
        print(f"注記: household が {len(case['household'])} か月分ある。書式は2か月分のため先頭2か月を使用")

    item_rows = fillmap["item_rows"]
    sheet = "家計収支表"
    filler = xlsxlib.XlsxFiller(vdir / "template.xlsx")

    values = {"fields": {
        "title_year": title_year(months[0].get("month", "")),
        "period_m1": month_period(months[0].get("month", "")),
        "period_m2": month_period(months[1].get("month", "")) if len(months) > 1 else "",
        "carryover_m1": months[0].get("carryover"),
    }}
    filler.apply(fillmap, values)

    # 正規名に無い費目は「その他」行へ。行の割当は全月で共通にする（列ズレ防止）
    for kind, rows_map, sonota_keys in (
        ("income", item_rows["income"], ["その他"]),
        ("expense", item_rows["expense"], ["その他1", "その他2"]),
    ):
        unmapped = []
        for h in months:
            for key in (h.get(kind) or {}):
                if key not in rows_map or key in sonota_keys:
                    if key not in unmapped:
                        unmapped.append(key)
        slot_of = {}
        for key in unmapped:
            if sonota_keys:
                slot = sonota_keys.pop(0)
                slot_of[key] = rows_map[slot]
                filler.write(sheet, f"B{slot_of[key]}", f"その他　（{key}）", f"kakei:{key}")
                print(f"注記: 費目「{key}」は正規名に無いため「その他」行（行{slot_of[key]}）に記入した")
            else:
                filler.warnings.append(f"費目「{key}」が書式に対応せず「その他」行も満杯。questions.md で扱いを確認")
        for col, h in zip(("C", "D"), months):
            for key, amount in (h.get(kind) or {}).items():
                if amount is None:
                    continue
                if key in rows_map and key not in slot_of:
                    filler.write(sheet, f"{col}{rows_map[key]}", amount, f"kakei:{key}")
                elif key in slot_of:
                    filler.write(sheet, f"{col}{slot_of[key]}", amount, f"kakei:{key}")

    filler.save(out)
    for w in filler.warnings:
        print(f"警告: {w}")
    return {"months": month_stats(months)}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--case", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--court", default="osaka")
    ap.add_argument("--generic", action="store_true")
    args = ap.parse_args()

    case = yaml.safe_load(open(args.case, encoding="utf-8"))
    vdir = None if args.generic else xlsxlib.resolve_form(args.court, "kakei-shushi")
    if vdir is None and not args.generic:
        print(xlsxlib.registration_guidance(args.court, "kakei-shushi", "家計収支表"))
        sys.exit(3)

    stats = build_from_template(case, vdir, args.output) if vdir else build_generic(case, args.output)
    Path(args.output + ".meta.json").write_text(
        json.dumps({"doc": "kakei-shushi", **stats}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"出力: {args.output}")
    for m in stats["months"]:
        print(f"  {m['month']}: 収入 {m['income']:,} / 支出 {m['expense']:,} / 収支 {m['balance']:+,}")


if __name__ == "__main__":
    main()
