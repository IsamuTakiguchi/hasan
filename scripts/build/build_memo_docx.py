#!/usr/bin/env python3
"""spec.yaml → 事務所内文書（聴取シート・面談メモ等）の docx を生成する汎用ビルダー。

裁判所書式ではない自由レイアウト文書用（fillmap 不要）。python-docx で新規生成する。

spec.yaml の形式:
  title: 聴取シート（破産申立て）
  subtitle: 乙山ストア事例 第2回面談          # 任意
  meta:                                      # 任意（表題下の左寄せ行）
    - "面談日：令和　年　月　日"
    - "出席者："
  sections:
    - heading: １　経歴（申立ての7年前から現在まで）
      note: 勤め先・自営・法人代表の別、期間、月収   # 任意（見出し下の補足・小さめ）
      items:
        - text: 現在の勤め先・月収
          check: true       # 先頭に □
          blank: true       # 末尾に記入用下線
        - text: "前職：株式会社乙山ストア（H22.6〜R7.2）で確認済み"
          done: true        # 先頭に ■（確認済み表示）
        - text: 自由記入欄
          lines: 3          # 記入用下線行を n 行追加
  footer: ※ 本シートは事務所内資料（提出書類ではない）   # 任意

使い方:
  python3 scripts/build/build_memo_docx.py --spec spec.yaml --output out.docx
"""
import argparse

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

BLANK = "＿" * 28
FONT = "ＭＳ 明朝"


def set_font(run, size=10.5, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)


def add_para(doc, text, size=10.5, bold=False, align=None, indent_mm=None, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if indent_mm:
        p.paragraph_format.left_indent = Mm(indent_mm)
    p.paragraph_format.space_after = Pt(space_after)
    set_font(p.add_run(text), size=size, bold=bold)
    return p


def build(spec, out):
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(20)

    add_para(doc, spec["title"], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    if spec.get("subtitle"):
        add_para(doc, spec["subtitle"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    for m in spec.get("meta", []) or []:
        add_para(doc, m, size=10.5, space_after=2)
    if spec.get("meta"):
        add_para(doc, "", space_after=2)

    for sec in spec.get("sections", []) or []:
        add_para(doc, sec["heading"], size=12, bold=True, space_after=3)
        if sec.get("note"):
            add_para(doc, f"（{sec['note']}）", size=9, indent_mm=4, space_after=3)
        for it in sec.get("items", []) or []:
            text = it["text"] if isinstance(it, dict) else str(it)
            it = it if isinstance(it, dict) else {}
            prefix = "□ " if it.get("check") else ("■ " if it.get("done") else "・")
            line = f"{prefix}{text}"
            if it.get("blank"):
                line += f"　{BLANK}"
            add_para(doc, line, size=10.5, indent_mm=5, space_after=3)
            for _ in range(int(it.get("lines", 0))):
                add_para(doc, "＿" * 40, size=10.5, indent_mm=9, space_after=3)
        add_para(doc, "", space_after=3)

    if spec.get("footer"):
        add_para(doc, spec["footer"], size=9, space_after=0)

    doc.save(out)
    n_items = sum(len(s.get("items", []) or []) for s in spec.get("sections", []) or [])
    print(f"出力: {out}（{len(spec.get('sections', []) or [])}節・{n_items}項目）")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    build(yaml.safe_load(open(args.spec, encoding="utf-8")), args.output)


if __name__ == "__main__":
    main()
