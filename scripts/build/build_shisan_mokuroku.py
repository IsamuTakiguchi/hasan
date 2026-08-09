#!/usr/bin/env python3
"""資産目録を case.yaml から生成する。

裁判所書式（courts/<court>/forms/shisan-mokuroku/）が登録済みなら fillmap 駆動、
未登録なら --generic で事務所内ドラフト様式を出力する。
<output>.meta.json に資産総額を書き出し、crosscheck_outputs.py が突合に使う。

使い方:
  python3 scripts/build/build_shisan_mokuroku.py --case cases/<id>/case.yaml \
      --output cases/<id>/output/資産目録_<姓>.xlsx [--court osaka] [--generic]
"""
import argparse
import json
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsxlib  # noqa: E402

import openpyxl  # noqa: E402


def collect_items(assets):
    """(区分, 名称・明細, 評価額, 備考) のリストと合計を返す。"""
    items = []

    def add(cat, name, value, note=""):
        items.append({"cat": cat, "name": name, "value": value, "note": note})

    if assets.get("cash") is not None:
        add("現金", "手持ち現金", assets["cash"])
    for d in assets.get("deposits", []) or []:
        name = f"{d.get('bank', '')}{d.get('branch', '') and ' ' + d['branch']} {d.get('account_type', '')} {d.get('account_no', '')}".strip()
        add("預貯金", name, d.get("balance"), d.get("as_of", "") and f"基準日 {d['as_of']}")
    for x in assets.get("insurance", []) or []:
        add("保険", f"{x.get('company', '')} {x.get('kind', '')}".strip(), x.get("surrender_value"), "解約返戻金見込")
    for x in assets.get("vehicles", []) or []:
        note = ""
        if x.get("loan_balance"):
            note = f"所有権留保ローン残 {x['loan_balance']:,}円"
        add("自動車", x.get("name", ""), x.get("value"), note)
    for x in assets.get("real_estate", []) or []:
        note = f"被担保債権残 {x['mortgage']:,}円" if x.get("mortgage") else ""
        add("不動産", f"{x.get('kind', '')} {x.get('address', '')}".strip(), x.get("value"), note)
    if assets.get("retirement_expectation") is not None:
        add("退職金", "退職金支給見込額", assets["retirement_expectation"])
    for x in assets.get("tsumitate", []) or []:
        add("積立金等", x.get("kind", ""), x.get("total"))
    for x in assets.get("shikikin", []) or []:
        add("敷金・保証金", x.get("property", ""), x.get("refund"),
            f"差入 {x['deposit']:,}円" if x.get("deposit") else "")
    for x in assets.get("loans_receivable", []) or []:
        add("貸付金等", x.get("debtor", ""), x.get("amount"),
            "" if x.get("recoverable") else "回収困難")
    for x in assets.get("upcoming", []) or []:
        add("取得見込財産", x.get("kind", ""), x.get("amount"))
    for x in assets.get("overpayment", []) or []:
        add("過払金等", x.get("against", ""), x.get("amount"))
    for x in assets.get("other", []) or []:
        add("その他", x.get("name", ""), x.get("value"), x.get("note", ""))

    total = sum(x["value"] for x in items if isinstance(x["value"], int))
    return items, total


def build_generic(case, out):
    meta = case.get("meta", {})
    items, total = collect_items(case.get("assets", {}) or {})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "資産目録"
    ws["A1"] = "資産目録"
    ws["A1"].font = xlsxlib.FONT_TITLE
    ws["A2"] = f"申立人: {meta.get('applicant', {}).get('name', '')}"
    ws["A2"].font = xlsxlib.FONT

    headers = ["区分", "名称・明細", "評価額(円)", "備考"]
    for i, (h, w) in enumerate(zip(headers, [10, 40, 14, 30]), 1):
        ws.cell(row=4, column=i, value=h)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    xlsxlib.style_header(ws, 4, len(headers))

    row = 5
    for it in items:
        ws.cell(row=row, column=1, value=it["cat"])
        ws.cell(row=row, column=2, value=it["name"])
        ws.cell(row=row, column=3, value=it["value"])
        ws.cell(row=row, column=4, value=it["note"])
        xlsxlib.style_data(ws, row, len(headers))
        ws[f"C{row}"].number_format = "#,##0"
        row += 1

    ws.cell(row=row, column=2, value="合計")
    ws.cell(row=row, column=3, value=f"=SUM(C5:C{row - 1})")
    xlsxlib.style_data(ws, row, len(headers))
    ws[f"C{row}"].number_format = "#,##0"
    ws[f"B{row}"].font = openpyxl.styles.Font(name="游ゴシック", size=10, bold=True)
    xlsxlib.draft_note(ws, row + 2)

    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"total_assets": total, "item_count": len(items)}


def wareki_md(s):
    """'R7.7.31' → '7月31日'。月日が取れなければそのまま。"""
    import re
    m = re.match(r"^[MTSHR]\d+\.(\d+)\.(\d+)$", s or "")
    return f"{m.group(1)}月{m.group(2)}日" if m else (s or "")


def build_from_template(case, vdir, out):
    """B1109（5シート・15区分）への記入。区分の【□有□無】は case.yaml で確認できるものだけ。"""
    fillmap = yaml.safe_load((vdir / "fillmap.yaml").read_text(encoding="utf-8"))
    assets = case.get("assets", {}) or {}
    menseki = case.get("menseki", {}) or {}
    f = {}
    unconfirmed = []

    def umu(field_id, key, data):
        """assets にキーが存在すれば 有/無 を返し、無ければ未確認として空欄。"""
        if key is not None and key not in assets:
            unconfirmed.append(field_id)
            return
        f[field_id] = "有" if data else "無"

    # 1 現金（5万円以上）
    cash = assets.get("cash")
    if "cash" in assets:
        f["cash_umu"] = "有" if (cash or 0) >= 50000 else "無"
        if (cash or 0) >= 50000:
            f["cash_amount"] = f"金額（{cash:,}円）"
            f["cash_keeper_dairinin"] = "申立代理人保管"
    else:
        unconfirmed.append("cash_umu")

    # 2 預貯金
    deposits = assets.get("deposits", []) or []
    umu("yochokin_umu", "deposits", deposits)
    futsu = [d for d in deposits if "定期" not in (d.get("account_type") or "")]
    teiki = [d for d in deposits if "定期" in (d.get("account_type") or "")]
    def dep_row(d):
        return {"bank": f"{d.get('bank', '')}({d.get('branch', '')})" if d.get("branch") else d.get("bank", ""),
                "account_no": d.get("account_no", ""), "balance": d.get("balance"),
                "as_of": wareki_md(d.get("as_of", "")), "kind": d.get("account_type", "")}
    if futsu:
        f["deposits_futsu"] = [dep_row(d) for d in futsu]
    if teiki:
        f["deposits_teiki"] = [dep_row(d) for d in teiki]

    # 3 保険（申立人名義）
    insurance = assets.get("insurance", []) or []
    umu("hoken_umu", "insurance", insurance)
    if insurance:
        rows = []
        total_unkaiyaku = 0
        for x in insurance:
            row = {"company": x.get("company", ""), "surrender_value": x.get("surrender_value"),
                   "kaiyaku": "□未解約"}
            if any(k in (x.get("kind") or "") for k in ("生命", "終身", "養老")):
                row["kind_check"] = "□生命保険"
            rows.append(row)
            if isinstance(x.get("surrender_value"), int):
                total_unkaiyaku += x["surrender_value"]
        f["hoken_honnin"] = rows
        f["hoken_surrender_total"] = total_unkaiyaku

    # 4 積立金 / 5 敷金 / 6 貸付金
    umu("tsumitate_umu", "tsumitate", assets.get("tsumitate"))
    if assets.get("tsumitate"):
        f["tsumitate_rows"] = [{"kind": x.get("kind", ""), "start": x.get("start", ""),
                                "total": x.get("total")} for x in assets["tsumitate"]]
    umu("shikikin_umu", "shikikin", assets.get("shikikin"))
    if assets.get("shikikin"):
        f["shikikin_rows"] = [{"property": x.get("property", ""), "start": x.get("start", ""),
                               "deposit": x.get("deposit"), "refund": x.get("refund"),
                               "arrears": x.get("arrears")} for x in assets["shikikin"]]
    umu("kashitsuke_umu", "loans_receivable", assets.get("loans_receivable"))
    if assets.get("loans_receivable"):
        f["kashitsuke_rows"] = [{"debtor": x.get("debtor", ""), "amount": x.get("amount"),
                                 "date": x.get("date", ""),
                                 "recoverable": "□有" if x.get("recoverable") else "□無",
                                 "reason": x.get("reason", "")} for x in assets["loans_receivable"]]

    # 7 退職金
    ret = assets.get("retirement_expectation")
    if "retirement_expectation" in assets:
        f["taishokukin_umu"] = "有" if (ret or 0) > 0 else "無"
        if (ret or 0) > 0:
            current = next((c for c in case.get("career", []) or []
                            if (c.get("period") or {}).get("end") == "現在"), None)
            if current:
                f["taishokukin_company"] = current.get("employer", "")
            f["taishokukin_amount"] = ret
            f["taishokukin_juryo_mijuryo"] = "受領未了"
    else:
        unconfirmed.append("taishokukin_umu")

    # 8 不動産
    re_list = assets.get("real_estate", []) or []
    umu("fudosan_umu", "real_estate", re_list)
    if re_list:
        blocks = []
        for x in re_list:
            b = {"address": x.get("address", ""), "owner_check": "□申立人", "mochibun_check": "□全部"}
            kind = x.get("kind", "")
            if "マンション" in kind:
                b["mansion_check"] = "□マンション"
            elif "土地" in kind:
                b["kind_check"] = "□土地"
            elif "建物" in kind or "戸建" in kind:
                b["kind_check"] = "□建物"
            blocks.append(b)
        f["fudosan_blocks"] = blocks

    # 9 自動車
    vehicles = assets.get("vehicles", []) or []
    umu("jidosha_umu", "vehicles", vehicles)
    if vehicles:
        blocks = []
        for x in vehicles:
            b = {"name": x.get("name", ""), "year": x.get("first_registration", ""),
                 "value": x.get("value")}
            if x.get("loan_balance"):
                b["ryuho_ari"] = "□有"
            else:
                b["ryuho_nashi"] = "□無"
            blocks.append(b)
        f["jidosha_blocks"] = blocks

    # 10 動産 / 11 その他財産（other は 11 に記入。動産該当は intake が note で区別する将来課題）
    other = assets.get("other", []) or []
    umu("dousan_umu", "other", [x for x in other if (x.get("value") or 0) >= 100000 and x.get("kind") == "動産"])
    umu("sonota_zaisan_umu", "other", other)
    if other:
        f["sonota_zaisan_rows"] = [{"name": x.get("name", ""), "value": x.get("value")} for x in other]

    # 12 処分財産 / 13 偏頗弁済
    disposal = menseki.get("asset_disposal")
    if disposal is not None:
        f["shobun_umu"] = "有" if disposal else "無"
        if disposal:
            f["shobun_rows"] = [{"item": x.get("item", ""), "date": x.get("sold", ""),
                                 "price": x.get("sold_price")} for x in disposal]
    henpa = (menseki.get("henpa") or {}).get("exists")
    if henpa is not None:
        f["henpa_umu"] = "有" if henpa else "無"
        if henpa:
            rows = [{"date": b.get("date", ""), "counterparty": b.get("note", "")[:20],
                     "amount": b.get("amount")}
                    for b in case.get("bank_analysis", []) or []
                    if b.get("flag") in ("受任通知後弁済", "偏頗弁済疑い") and b.get("resolved")]
            if rows:
                f["henpa_rows"] = rows

    # 14 取得見込 / 15 過払金
    umu("kinjitsu_umu", "upcoming", assets.get("upcoming"))
    if assets.get("upcoming"):
        f["kinjitsu_rows"] = [{"kind": x.get("kind", ""), "counterparty": x.get("counterparty", ""),
                               "amount": x.get("amount")} for x in assets["upcoming"]]
    op = assets.get("overpayment", []) or []
    umu("kabarai_umu", "overpayment", op)
    if op:
        f["kabarai_rows"] = [{"against": x.get("against", ""), "amount": x.get("amount")} for x in op]

    filler = xlsxlib.XlsxFiller(vdir / "template.xlsx")
    filler.apply(fillmap, {"fields": f})
    filler.save(out)
    for w in filler.warnings:
        print(f"警告: {w}")
    if unconfirmed:
        print(f"未確認（有無チェックを空欄にした）: {', '.join(unconfirmed)} → questions.md で確認")

    items, total = collect_items(assets)
    return {"total_assets": total, "item_count": len(items)}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--case", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--court", default="osaka")
    ap.add_argument("--generic", action="store_true")
    args = ap.parse_args()

    case = yaml.safe_load(open(args.case, encoding="utf-8"))
    vdir = None if args.generic else xlsxlib.resolve_form(
        args.court, "shisan-mokuroku", proc=xlsxlib.proc_of_case(case))
    if vdir is None and not args.generic:
        print(xlsxlib.registration_guidance(args.court, "shisan-mokuroku", "資産目録"))
        sys.exit(3)

    stats = build_from_template(case, vdir, args.output) if vdir else build_generic(case, args.output)
    Path(args.output + ".meta.json").write_text(
        json.dumps({"doc": "shisan-mokuroku", **stats}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"出力: {args.output}")
    print(f"資産 {stats['item_count']} 項目 / 総額 {stats['total_assets']:,} 円")


if __name__ == "__main__":
    main()
